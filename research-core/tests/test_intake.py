from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from trading_research_core.errors import CoreError
from trading_research_core.intake import (
    get_evidence,
    intake_mt5_excel,
    list_registry,
    verify_raw_snapshot,
)
from trading_research_core.mt5_excel import REQUIRED_DEAL_HEADERS, file_sha256


def _populate_supported_report(sheet: object, *, close_comment: str = "Close") -> None:
    for row, (key, value) in enumerate(
        (
            ("Expert", "Example EA"),
            ("Symbol", "EURUSD"),
            ("Period", "H1"),
            ("Currency", "USD"),
            ("Initial Deposit", "1000"),
            ("Leverage", "1:100"),
        ),
        start=1,
    ):
        sheet.cell(row=row, column=1, value=f"{key}:")  # type: ignore[attr-defined]
        sheet.cell(row=row, column=4, value=value)  # type: ignore[attr-defined]

    deals_row = 8
    sheet.cell(row=deals_row, column=1, value="Deals")  # type: ignore[attr-defined]
    for column, header in enumerate(REQUIRED_DEAL_HEADERS, start=1):
        sheet.cell(row=deals_row + 1, column=column, value=header)  # type: ignore[attr-defined]
    for column, value in enumerate(["2026.01.01 00:00:00", "1", None, "Balance", None, None, None, None, 0, 0, 0, 1000, "Opening balance"], start=1):
        sheet.cell(row=deals_row + 2, column=column, value=value)  # type: ignore[attr-defined]
    for column, value in enumerate(["2026.01.02 00:00:00", "2", "EURUSD", "Buy", "Out", 0.1, 1.10000, "10", -1, 0, 11, 1010, close_comment], start=1):
        sheet.cell(row=deals_row + 3, column=column, value=value)  # type: ignore[attr-defined]
    sheet.cell(row=deals_row + 3, column=7).number_format = "0.00000"  # type: ignore[attr-defined]


def _write_supported_report(path: Path, *, close_comment: str = "Close") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    _populate_supported_report(sheet, close_comment=close_comment)
    workbook.save(path)
    workbook.close()


def test_intake_preserves_bytes_and_reuses_identical_source(tmp_path: Path) -> None:
    source = tmp_path / "mt5-report.xlsx"
    workspace = tmp_path / "workspace"
    _write_supported_report(source)
    source_bytes = source.read_bytes()
    source_hash = file_sha256(source)

    first = intake_mt5_excel(workspace, str(source))
    receipt = first["intake_receipt"]
    assert isinstance(receipt, dict)
    snapshot = Path(str(receipt["raw_snapshot_path"]))
    assert first["intake_status"] == "SNAPSHOT_CREATED"
    assert source.read_bytes() == source_bytes
    assert snapshot.read_bytes() == source_bytes
    assert snapshot == workspace / "raw" / source_hash / "source.xlsx"
    assert verify_raw_snapshot(workspace, str(first["dataset_ref"]))["verified"] is True
    assert receipt["supplied_facts"] == {"symbol": "EURUSD", "currency": "USD", "period": "H1", "initial_deposit": "1000", "leverage": "1:100"}
    assert receipt["observed_price_scales"] == [5]
    assert receipt["artifacts"]["events"] == f"{first['dataset_ref']}:events"

    second = intake_mt5_excel(workspace, str(source))
    assert second["intake_status"] == "REUSED_IDENTICAL_SOURCE"
    assert second["dataset_ref"] == first["dataset_ref"]
    registry = list_registry(workspace)
    assert registry["registry_schema_version"] == "1.0"
    assert len(registry["entries"]) == 1
    assert get_evidence(workspace, str(first["dataset_ref"]))["source_sha256"] == source_hash
    registry_json = json.loads((workspace / "registry" / "datasets.json").read_text(encoding="utf-8"))
    assert registry_json["entries"][0]["raw_snapshot_status"] == "VERIFIED"


def test_invalid_source_does_not_create_managed_intake_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "not-an-mt5-report.xlsx"
    workspace = tmp_path / "workspace"
    workbook = Workbook()
    workbook.active["A1"] = "This is not an MT5 report"
    workbook.save(source)
    workbook.close()

    with pytest.raises(CoreError, match="Missing required MT5 Settings values") as error:
        intake_mt5_excel(workspace, str(source))

    assert error.value.code == "E_SOURCE_INVALID"
    assert not (workspace / "raw").exists()
    assert not (workspace / "registry").exists()


def test_ambiguous_mt5_deals_worksheets_are_rejected_without_intake_writes(tmp_path: Path) -> None:
    source = tmp_path / "ambiguous-report.xlsx"
    workspace = tmp_path / "workspace"
    workbook = Workbook()
    _populate_supported_report(workbook.active)
    _populate_supported_report(workbook.create_sheet("Second report"))
    workbook.save(source)
    workbook.close()

    with pytest.raises(CoreError, match="more than one possible MT5 Deals worksheet") as error:
        intake_mt5_excel(workspace, str(source))

    assert error.value.code == "E_SOURCE_AMBIGUOUS"
    assert error.value.details["worksheet_names"] == ["Sheet", "Second report"]
    assert not (workspace / "raw").exists()
    assert not (workspace / "datasets").exists()
    assert not (workspace / "registry").exists()


def test_byte_modified_valid_report_creates_a_distinct_dataset(tmp_path: Path) -> None:
    first_source = tmp_path / "first.xlsx"
    modified_source = tmp_path / "modified.xlsx"
    workspace = tmp_path / "workspace"
    _write_supported_report(first_source, close_comment="Close")
    _write_supported_report(modified_source, close_comment="Changed source fixture")

    first = intake_mt5_excel(workspace, str(first_source))
    modified = intake_mt5_excel(workspace, str(modified_source))

    assert first["source"]["sha256"] != modified["source"]["sha256"]  # type: ignore[index]
    assert first["dataset_ref"] != modified["dataset_ref"]
    assert modified["intake_status"] == "SNAPSHOT_CREATED"
    assert len(list_registry(workspace)["entries"]) == 2


def test_tampered_managed_snapshot_is_detected(tmp_path: Path) -> None:
    source = tmp_path / "mt5-report.xlsx"
    workspace = tmp_path / "workspace"
    _write_supported_report(source)
    source_bytes = source.read_bytes()
    intake = intake_mt5_excel(workspace, str(source))
    receipt = intake["intake_receipt"]
    assert isinstance(receipt, dict)
    Path(str(receipt["raw_snapshot_path"])).write_bytes(b"tampered test snapshot")

    verification = verify_raw_snapshot(workspace, str(intake["dataset_ref"]))
    assert verification["verified"] is False
    assert verification["observed_sha256"] != verification["expected_sha256"]
    assert source.read_bytes() == source_bytes
    with pytest.raises(CoreError, match="Existing managed raw snapshot hash does not match") as error:
        intake_mt5_excel(workspace, str(source))
    assert error.value.code == "E_RAW_SNAPSHOT_MISMATCH"
