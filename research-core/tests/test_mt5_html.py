"""MT5 Strategy Tester HTML reports: same canonical result as `.xlsx`, plus a totals check."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from trading_research_core.errors import CoreError
from trading_research_core.intake import get_evidence, intake_mt5_report, list_registry, verify_raw_snapshot
from trading_research_core.mt5_excel import REQUIRED_DEAL_HEADERS, import_mt5_excel
from trading_research_core.mt5_html import import_mt5_html, read_html_rows
from trading_research_core.mt5_report_summary import read_report_summary
from trading_research_core.worker import Worker

SETTINGS = [("Expert", "Example EA"), ("Symbol", "EURUSD"), ("Period", "H1 (2026.01.01 - 2026.01.31)"), ("Company", "Demo Broker"), ("Currency", "USD"), ("Initial Deposit", "10 000.00"), ("Leverage", "1:100")]
INPUTS = ["InpLot=0.1", "InpMode=2"]
RESULTS = [("Total Net Profit", "1 234.50"), ("Profit Factor", "1.85"), ("Total Trades", "2"), ("Equity Drawdown Relative", "3.21% (321.00)")]
DEALS = [
    ["2026.01.01 00:00:00", "1", None, "balance", None, None, None, None, "0.00", "0.00", "10 000.00", "10 000.00", None],
    ["2026.01.02 10:00:00", "2", "EURUSD", "buy", "in", "0.10", "1.10000", "2", "-0.50", "0.00", "0.00", "9 999.50", "open"],
    ["2026.01.03 10:00:00", "3", "EURUSD", "sell", "out", "0.10", "1.12345", "3", "-0.50", "-1.00", "1 236.50", "11 234.50", "close"],
]
TOTALS = [None] * 8 + ["-1.00", "-1.00", "1 236.50", "11 234.50", None]


def _html(deals: list[list[str | None]] = DEALS, totals: list[str | None] | None = TOTALS) -> str:
    def row(cells: list[tuple[str | None, int]]) -> str:
        return "<tr align=\"right\">" + "".join(f"<td nowrap colspan=\"{span}\"><b>{value or ''}</b></td>" for value, span in cells) + "</tr>\n"

    body = row([("Strategy Tester Report", 13)]) + row([("Settings", 13)])
    for label, value in SETTINGS[:3]:
        body += row([(f"{label}:", 3), (value, 10)])
    body += row([("Inputs:", 3), (INPUTS[0], 10)]) + row([(None, 3), (INPUTS[1], 10)])
    for label, value in SETTINGS[3:]:
        body += row([(f"{label}:", 3), (value, 10)])
    body += row([("Results", 13)])
    for label, value in RESULTS:
        body += row([(f"{label}:", 3), (value, 10)])
    body += row([("Orders", 13)]) + row([("Deals", 13)]) + "<tr>" + "".join(f"<th>{header}</th>" for header in REQUIRED_DEAL_HEADERS) + "</tr>\n"
    for deal in deals:
        body += "<tr>" + "".join(f"<td>{value if value is not None else ''}</td>" for value in deal) + "</tr>\n"
    if totals is not None:
        body += "<tr>" + "".join(f"<td>{value if value is not None else ''}</td>" for value in totals) + "</tr>\n"
    return f"<!DOCTYPE html><html><head><title>Strategy Tester Report</title></head><body><table>{body}</table><img src=\"chart.png\"></body></html>"


def _write_html(path: Path, text: str | None = None, encoding: str = "utf-16") -> Path:
    content = text if text is not None else _html()
    path.write_bytes(("﻿" + content).encode("utf-16-le") if encoding == "utf-16" else content.encode("utf-8"))
    return path


def _write_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    number = lambda value: None if value is None else float(value.replace(" ", "")) if any(ch.isdigit() for ch in value) and value.replace(" ", "").replace(".", "").lstrip("-").isdigit() else value
    for index, (label, value) in enumerate(SETTINGS, start=1):
        sheet.cell(row=index, column=1, value=f"{label}:")
        sheet.cell(row=index, column=4, value=value.replace(" ", "") if label == "Initial Deposit" else value)
    sheet.cell(row=10, column=1, value="Deals")
    for column, header in enumerate(REQUIRED_DEAL_HEADERS, start=1):
        sheet.cell(row=11, column=column, value=header)
    for offset, deal in enumerate(DEALS):
        for column, value in enumerate(deal, start=1):
            cell = sheet.cell(row=12 + offset, column=column, value=value if column in (1, 2, 3, 4, 5, 6, 8, 13) else number(value))
            if column == 7 and value is not None:
                cell.number_format = "0.00000"
    workbook.save(path)
    workbook.close()
    return path


def test_html_matches_the_excel_adapter(tmp_path: Path) -> None:
    html = import_mt5_html(str(_write_html(tmp_path / "report.htm")))
    xlsx = import_mt5_excel(str(_write_xlsx(tmp_path / "report.xlsx")))
    strip = lambda events: [{key: value for key, value in event.items() if key != "source_artifact_sha256"} for event in events]
    assert _normalised(strip(html["events"])) == _normalised(strip(xlsx["events"]))  # same values; HTML text vs Excel numbers
    assert html["settings"]["Initial Deposit"] == "10000.00" and html["settings"]["Period"] == "H1 (2026.01.01 - 2026.01.31)"
    assert html["events"][2]["source_profit"] == "1236.50" and html["events"][2]["source_price_scale"] == 5
    assert html["source_checks"] == ["HTML_DEALS_TOTALS_MATCH"] and html["source"]["worksheet_name"] == "HTML report"


def _normalised(events: list[dict[str, object]]) -> list[dict[str, object]]:
    from decimal import Decimal

    def value(item: object) -> object:
        try:
            return Decimal(str(item)).normalize() if isinstance(item, (str, float)) else item
        except ArithmeticError:
            return item
    return [{key: value(item) for key, item in event.items()} for event in events]


def test_utf8_and_missing_totals_are_accepted(tmp_path: Path) -> None:
    result = import_mt5_html(str(_write_html(tmp_path / "report.html", _html(totals=None), encoding="utf-8")))
    assert len(result["events"]) == 3 and result["source_checks"] == ["HTML_DEALS_TOTALS_ABSENT"]


def test_totals_mismatch_blocks(tmp_path: Path) -> None:
    wrong = TOTALS[:10] + ["999.00"] + TOTALS[11:]
    with pytest.raises(CoreError) as error:
        import_mt5_html(str(_write_html(tmp_path / "report.htm", _html(totals=wrong))))
    assert error.value.code == "E_SOURCE_INVALID" and error.value.details["differing"] == ["profit"]


@pytest.mark.parametrize("text,code", [("<html><body><p>hello</p></body></html>", "E_SOURCE_INVALID"), ("<html><table><tr><td>Settings</td></tr></table></html>", "E_SOURCE_INVALID")])
def test_non_reports_are_rejected(tmp_path: Path, text: str, code: str) -> None:
    with pytest.raises(CoreError) as error:
        import_mt5_html(str(_write_html(tmp_path / "page.htm", text)))
    assert error.value.code == code


def test_colspans_expand_to_excel_columns(tmp_path: Path) -> None:
    rows = read_html_rows(_write_html(tmp_path / "report.htm"))
    expert = next(row for row in rows if row[0].value == "Expert:")
    assert [cell.value for cell in expert[:4]] == ["Expert:", None, None, "Example EA"] and len(expert) == 13


def test_intake_snapshots_html_and_registers_its_adapter(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    source = _write_html(tmp_path / "report.html")
    first = intake_mt5_report(workspace, str(source))
    receipt = first["intake_receipt"]
    assert Path(str(receipt["raw_snapshot_path"])).name == "source.htm" and Path(str(receipt["raw_snapshot_path"])).read_bytes() == source.read_bytes()
    assert receipt["adapter"]["adapter_id"] == "mt5-strategy-tester-html" and receipt["source_checks"] == ["HTML_DEALS_TOTALS_MATCH"]
    assert receipt["supplied_facts"]["initial_deposit"] == "10000.00"
    assert verify_raw_snapshot(workspace, str(first["dataset_ref"]))["verified"] is True
    assert intake_mt5_report(workspace, str(source))["intake_status"] == "REUSED_IDENTICAL_SOURCE"
    assert [entry["dataset_ref"] for entry in list_registry(workspace)["entries"]] == [first["dataset_ref"]]
    summary = read_report_summary(Path(str(get_evidence(workspace, str(first["dataset_ref"]))["raw_snapshot_path"])))
    assert summary["inputs"] == {"InpLot": "0.1", "InpMode": "2"}
    assert (summary["metrics"]["net_profit"], summary["metrics"]["equity_drawdown_pct"], summary["timeframe"], summary["start"]) == ("1234.50", "3.21", "H1", "2026.01.01")


def test_worker_intakes_html_through_both_methods(tmp_path: Path) -> None:
    worker = Worker(tmp_path / "workspace")
    source = _write_html(tmp_path / "report.htm")
    methods = worker.dispatch({"method": "core.capabilities", "params": {}})["methods"]
    assert "dataset.intake_mt5_report" in methods
    first = worker.dispatch({"method": "dataset.intake_mt5_report", "params": {"source_path": str(source)}})
    again = worker.dispatch({"method": "dataset.intake_mt5_excel", "params": {"source_path": str(source)}})
    assert first["dataset_ref"] == again["dataset_ref"] and json.dumps(first["event_count"]) == "3"
