"""Read the observed MT5 Strategy Tester Excel Deals export faithfully."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

from .errors import CoreError


REQUIRED_DEAL_HEADERS = (
    "Time", "Deal", "Symbol", "Type", "Direction", "Volume", "Price", "Order",
    "Commission", "Swap", "Profit", "Balance", "Comment",
)


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def import_mt5_excel(path_text: str) -> dict[str, object]:
    """Return source-preserving canonical events from a supported .xlsx report."""

    source_path = Path(path_text).expanduser().resolve()
    if not source_path.is_file():
        raise CoreError("E_SOURCE_NOT_FOUND", "Source file does not exist.", details={"path": str(source_path)})
    if source_path.suffix.lower() != ".xlsx":
        raise CoreError("E_SOURCE_UNSUPPORTED", "M0 supports MT5 Strategy Tester .xlsx exports only.")
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise CoreError("E_DEPENDENCY_MISSING", "openpyxl is required for MT5 Excel import.") from error

    workbook = load_workbook(source_path, read_only=True, data_only=False)
    try:
        if not workbook.worksheets:
            raise CoreError("E_SOURCE_INVALID", "Workbook has no worksheets.")
        deal_sheets = [sheet for sheet in workbook.worksheets if _has_deals_marker(sheet)]
        if len(deal_sheets) > 1:
            raise CoreError(
                "E_SOURCE_AMBIGUOUS",
                "Workbook contains more than one possible MT5 Deals worksheet.",
                details={"worksheet_names": [sheet.title for sheet in deal_sheets]},
            )
        sheet = deal_sheets[0] if deal_sheets else workbook.worksheets[0]
        rows = list(sheet.iter_rows())
        settings = _extract_settings(rows)
        events = _extract_events(rows, file_sha256(source_path))
        worksheet_name = sheet.title
    finally:
        workbook.close()

    return {
        "source": {
            "filename": source_path.name,
            "sha256": file_sha256(source_path),
            "byte_count": source_path.stat().st_size,
            "worksheet_name": worksheet_name,
        },
        "settings": settings,
        "events": events,
    }


def _has_deals_marker(sheet: Any) -> bool:
    return any(row[0].value == "Deals" for row in sheet.iter_rows())


def _extract_settings(rows: list[tuple[Any, ...]]) -> dict[str, str]:
    settings: dict[str, str] = {}
    for row in rows:
        first = row[0].value
        if first == "Results":
            break
        if isinstance(first, str) and first.endswith(":") and len(row) > 3 and row[3].value is not None:
            settings[first[:-1]] = str(row[3].value)
    required = ("Expert", "Symbol", "Period", "Currency", "Initial Deposit", "Leverage")
    missing = [field for field in required if field not in settings]
    if missing:
        raise CoreError("E_SOURCE_INVALID", "Missing required MT5 Settings values.", details={"missing": missing})
    return settings


def _extract_events(rows: list[tuple[Any, ...]], artifact_sha256: str) -> list[dict[str, object]]:
    deals_index = next((index for index, row in enumerate(rows) if row[0].value == "Deals"), None)
    if deals_index is None or deals_index + 1 >= len(rows):
        raise CoreError("E_SOURCE_INVALID", "Workbook does not contain a Deals table.")
    headers = tuple(cell.value for cell in rows[deals_index + 1][:len(REQUIRED_DEAL_HEADERS)])
    if headers != REQUIRED_DEAL_HEADERS:
        raise CoreError("E_SOURCE_UNSUPPORTED", "Deals headers do not match the supported MT5 layout.", details={"received": list(headers)})

    events: list[dict[str, object]] = []
    for sequence, row in enumerate(rows[deals_index + 2 :], start=1):
        values = [cell.value for cell in row[:len(REQUIRED_DEAL_HEADERS)]]
        if not any(value is not None for value in values):
            continue
        if values[0] is None:
            raise CoreError("E_SOURCE_INVALID", "Deals row is missing Time.", details={"source_sequence": sequence})
        deal_type = _text(values[3], "Type", sequence).lower()
        direction = _optional_text(values[4])
        event_type, side = _event_kind(deal_type, direction, sequence)
        event = {
            "source_artifact_sha256": artifact_sha256,
            "source_sequence": sequence,
            "source_deal_id": _text(values[1], "Deal", sequence),
            "source_order_id": _optional_text(values[7]),
            "source_timestamp_text": str(values[0]),
            "source_timestamp": _timestamp(values[0], sequence),
            "event_type": event_type,
            "side": side,
            "symbol": _optional_text(values[2]),
            "volume": _decimal(values[5], "Volume", sequence, required=event_type != "OPENING_BALANCE"),
            "price": _decimal(values[6], "Price", sequence, required=event_type != "OPENING_BALANCE"),
            "source_price_scale": _price_scale(row[6].number_format) if values[6] is not None else None,
            "source_commission": _decimal(values[8], "Commission", sequence, required=True),
            "source_swap": _decimal(values[9], "Swap", sequence, required=True),
            "source_profit": _decimal(values[10], "Profit", sequence, required=True),
            "reported_balance": _decimal(values[11], "Balance", sequence, required=True),
            "comment": _optional_text(values[12]),
        }
        if event_type != "OPENING_BALANCE" and event["symbol"] is None:
            raise CoreError("E_SOURCE_INVALID", "Trading deal is missing Symbol.", details={"source_sequence": sequence})
        events.append(event)
    if not events:
        raise CoreError("E_SOURCE_INVALID", "Deals table contains no records.")
    if events[0]["event_type"] != "OPENING_BALANCE" or sum(event["event_type"] == "OPENING_BALANCE" for event in events) != 1:
        raise CoreError("E_SOURCE_UNSUPPORTED", "M0 requires exactly one opening balance as the first Deals row.")
    return events


def _event_kind(deal_type: str, direction: str | None, sequence: int) -> tuple[str, str | None]:
    normalized_direction = direction.lower() if direction is not None else None
    if deal_type == "balance":
        if normalized_direction is not None:
            raise CoreError("E_SOURCE_UNSUPPORTED", "Balance row has an unexpected Direction.", details={"source_sequence": sequence})
        return "OPENING_BALANCE", None
    if deal_type not in {"buy", "sell"} or normalized_direction not in {"in", "out"}:
        raise CoreError("E_SOURCE_UNSUPPORTED", "Unsupported MT5 deal semantics in M0.", details={"source_sequence": sequence, "type": deal_type, "direction": direction})
    return ("POSITION_OPEN" if normalized_direction == "in" else "POSITION_CLOSE"), deal_type.upper()


def _text(value: object, field: str, sequence: int) -> str:
    text = _optional_text(value)
    if text is None:
        raise CoreError("E_SOURCE_INVALID", f"Deals row is missing {field}.", details={"source_sequence": sequence})
    return text


def _optional_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _decimal(value: object, field: str, sequence: int, *, required: bool) -> str | None:
    if value is None:
        if required:
            raise CoreError("E_SOURCE_INVALID", f"Deals row is missing {field}.", details={"source_sequence": sequence})
        return None
    try:
        return str(Decimal(str(value)))
    except (InvalidOperation, ValueError) as error:
        raise CoreError("E_SOURCE_INVALID", f"Deals row has invalid {field}.", details={"source_sequence": sequence}) from error


def _timestamp(value: object, sequence: int) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    try:
        return datetime.strptime(str(value), "%Y.%m.%d %H:%M:%S").strftime("%Y-%m-%dT%H:%M:%S")
    except ValueError as error:
        raise CoreError("E_SOURCE_INVALID", "Unsupported MT5 timestamp.", details={"source_sequence": sequence, "value": str(value)}) from error


def _price_scale(number_format: str) -> int | None:
    matches = re.findall(r"[0#]+\.([0#]+)", number_format)
    return max((len(match) for match in matches), default=None)
