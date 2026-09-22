"""Read-only parser for the observed MT5 Strategy Tester Excel Deals layout."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


REQUIRED_DEAL_HEADERS = (
    "Time",
    "Deal",
    "Symbol",
    "Type",
    "Direction",
    "Volume",
    "Price",
    "Order",
    "Commission",
    "Swap",
    "Profit",
    "Balance",
    "Comment",
)


class ImportValidationError(ValueError):
    """Raised when a workbook is not the supported MT5 Deals report layout."""


@dataclass(frozen=True)
class SourceArtifact:
    """Immutable identity and provenance for one imported source file."""

    path: Path
    sha256: str
    byte_count: int
    worksheet_name: str


@dataclass(frozen=True)
class Mt5Deal:
    """One source Deals row, with source-signed cash fields."""

    timestamp_text: str
    timestamp: datetime
    deal_id: str
    symbol: str | None
    deal_type: str
    direction: str | None
    volume: Decimal | None
    price: Decimal | None
    source_price_scale: int | None
    order_id: str | None
    commission: Decimal
    swap: Decimal
    profit: Decimal
    balance: Decimal
    comment: str | None

    @property
    def net_cash_change(self) -> Decimal:
        """Source-signed realised cash movement for this deal row."""

        return self.profit + self.commission + self.swap


@dataclass(frozen=True)
class Mt5ExcelImport:
    """Parsed source records and report-level settings for one workbook."""

    artifact: SourceArtifact
    settings: dict[str, str]
    deals: tuple[Mt5Deal, ...]

    @property
    def opening_balance(self) -> Decimal:
        balances = [deal.balance for deal in self.deals if deal.deal_type == "balance"]
        if not balances:
            raise ImportValidationError("Deals table has no opening balance row.")
        return balances[0]

    @property
    def trading_deals(self) -> tuple[Mt5Deal, ...]:
        return tuple(deal for deal in self.deals if deal.deal_type != "balance")

    @property
    def net_trading_cash_change(self) -> Decimal:
        return sum((deal.net_cash_change for deal in self.trading_deals), Decimal("0"))


def import_mt5_excel_report(path: str | Path) -> Mt5ExcelImport:
    """Parse the supported MT5 Strategy Tester Excel report without modifying it."""

    source_path = Path(path).expanduser().resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    if source_path.suffix.lower() != ".xlsx":
        raise ImportValidationError("Only .xlsx MT5 Strategy Tester reports are supported in Step 3.")

    workbook = load_workbook(source_path, read_only=True, data_only=False)
    try:
        if not workbook.worksheets:
            raise ImportValidationError("Workbook has no worksheets.")
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows())
        settings = _extract_settings(rows)
        deals = _extract_deals(rows)
        worksheet_name = sheet.title
    finally:
        workbook.close()

    artifact = SourceArtifact(
        path=source_path,
        sha256=_file_sha256(source_path),
        byte_count=source_path.stat().st_size,
        worksheet_name=worksheet_name,
    )
    return Mt5ExcelImport(artifact=artifact, settings=settings, deals=tuple(deals))


def _extract_settings(rows: list[tuple[Any, ...]]) -> dict[str, str]:
    settings: dict[str, str] = {}
    for row in rows:
        first = row[0].value
        if first == "Results":
            break
        if isinstance(first, str) and first.endswith(":") and len(row) > 3:
            value = row[3].value
            if value is not None:
                settings[first[:-1]] = str(value)
    required = ("Expert", "Symbol", "Period", "Currency", "Initial Deposit", "Leverage")
    missing = [key for key in required if key not in settings]
    if missing:
        raise ImportValidationError(f"Missing required Settings values: {', '.join(missing)}")
    return settings


def _extract_deals(rows: list[tuple[Any, ...]]) -> list[Mt5Deal]:
    deals_index = next((index for index, row in enumerate(rows) if row[0].value == "Deals"), None)
    if deals_index is None:
        raise ImportValidationError("Workbook does not contain a Deals section.")
    header_index = deals_index + 1
    headers = tuple(cell.value for cell in rows[header_index][: len(REQUIRED_DEAL_HEADERS)])
    if headers != REQUIRED_DEAL_HEADERS:
        raise ImportValidationError(
            "Unsupported Deals headers. Expected "
            f"{REQUIRED_DEAL_HEADERS!r}, received {headers!r}."
        )

    parsed: list[Mt5Deal] = []
    for row in rows[header_index + 1 :]:
        values = [cell.value for cell in row[: len(REQUIRED_DEAL_HEADERS)]]
        if not any(value is not None for value in values):
            continue
        if values[0] is None:
            raise ImportValidationError("Deals row is missing Time.")
        parsed.append(
            Mt5Deal(
                timestamp_text=str(values[0]),
                timestamp=_parse_timestamp(values[0]),
                deal_id=_required_text(values[1], "Deal"),
                symbol=_optional_text(values[2]),
                deal_type=_required_text(values[3], "Type").lower(),
                direction=_optional_text(values[4]),
                volume=_optional_decimal(values[5]),
                price=_optional_decimal(values[6]),
                source_price_scale=_price_scale(row[6].number_format) if values[6] is not None else None,
                order_id=_optional_text(values[7]),
                commission=_required_decimal(values[8], "Commission"),
                swap=_required_decimal(values[9], "Swap"),
                profit=_required_decimal(values[10], "Profit"),
                balance=_required_decimal(values[11], "Balance"),
                comment=_optional_text(values[12]),
            )
        )
    if not parsed:
        raise ImportValidationError("Deals section contains no records.")
    return parsed


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _parse_timestamp(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value), "%Y.%m.%d %H:%M:%S")
    except ValueError as error:
        raise ImportValidationError(f"Unsupported MT5 timestamp: {value!r}") from error


def _required_text(value: Any, field: str) -> str:
    text = _optional_text(value)
    if text is None:
        raise ImportValidationError(f"Deals row is missing {field}.")
    return text


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _required_decimal(value: Any, field: str) -> Decimal:
    parsed = _optional_decimal(value)
    if parsed is None:
        raise ImportValidationError(f"Deals row is missing {field}.")
    return parsed


def _optional_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except Exception as error:  # Decimal exposes multiple exception subclasses.
        raise ImportValidationError(f"Unsupported decimal value: {value!r}") from error


def _price_scale(number_format: str) -> int | None:
    matches = re.findall(r"[0#]+\.([0#]+)", number_format)
    return max((len(match) for match in matches), default=None)


def _main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Read an MT5 Strategy Tester Excel Deals report.")
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    report = import_mt5_excel_report(args.report)
    print(f"artifact_sha256={report.artifact.sha256}")
    print(f"symbol={report.settings['Symbol']}")
    print(f"source_rows={len(report.deals)}")
    print(f"trading_deals={len(report.trading_deals)}")
    print(f"opening_balance={report.opening_balance}")
    print(f"net_trading_cash_change={report.net_trading_cash_change}")


if __name__ == "__main__":
    _main()
