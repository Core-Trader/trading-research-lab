from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from trading_research_lab.event_normalization import (
    CanonicalEventType,
    NormalizationError,
    normalize_mt5_import,
)
from trading_research_lab.mt5_excel import import_mt5_excel_report

from test_mt5_excel import write_report


class EventNormalizationTests(unittest.TestCase):
    def test_maps_supported_rows_without_pairing_trades(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "USDJPY.xlsx"
            write_report(path)
            imported = import_mt5_excel_report(path)

            result = normalize_mt5_import(
                imported, broker_time_profile_id="roboforex-eet-eest-v0.1"
            )

        self.assertEqual(len(result.events), len(imported.deals))
        self.assertEqual(
            [event.event_type for event in result.events],
            [
                CanonicalEventType.OPENING_BALANCE,
                CanonicalEventType.POSITION_OPEN,
                CanonicalEventType.POSITION_CLOSE,
            ],
        )
        self.assertEqual(result.opening_balance_event.reported_balance, Decimal("15000"))
        self.assertIsNone(result.opening_balance_event.source_net_cash_change)
        self.assertEqual(result.events[1].side, "BUY")
        self.assertEqual(result.events[2].side, "SELL")
        self.assertEqual(result.events[2].source_net_cash_change, Decimal("0.50"))
        self.assertEqual(result.events[1].source_price_scale, 3)
        self.assertEqual(result.events[1].broker_time_profile_id, "roboforex-eet-eest-v0.1")

    def test_event_ids_are_deterministic_for_the_same_import(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "EURUSD.xlsx"
            write_report(path, price=1.08543, price_format="0.00000")
            imported = import_mt5_excel_report(path)

            first = normalize_mt5_import(
                imported, broker_time_profile_id="roboforex-eet-eest-v0.1"
            )
            second = normalize_mt5_import(
                imported, broker_time_profile_id="roboforex-eet-eest-v0.1"

            )

        self.assertEqual(
            [event.event_id for event in first.events],
            [event.event_id for event in second.events],
        )

    def test_rejects_a_trade_with_an_unsupported_direction(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "unsupported-direction.xlsx"
            write_report(path)
            workbook = load_workbook(path)
            sheet = workbook.active
            deals_row = next(cell.row for cell in sheet["A"] if cell.value == "Deals")
            sheet.cell(row=deals_row + 3, column=5, value="inout")
            workbook.save(path)
            imported = import_mt5_excel_report(path)

            with self.assertRaises(NormalizationError):
                normalize_mt5_import(
                    imported, broker_time_profile_id="roboforex-eet-eest-v0.1"
                )

    def test_rejects_a_later_balance_adjustment(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "later-balance.xlsx"
            write_report(path)
            workbook = load_workbook(path)
            sheet = workbook.active
            sheet.append(
                [
                    "2025.01.11 00:00:00",
                    4,
                    None,
                    "balance",
                    None,
                    None,
                    None,
                    None,
                    0,
                    0,
                    100,
                    15100.50,
                    "deposit",
                ]
            )
            workbook.save(path)
            imported = import_mt5_excel_report(path)

            with self.assertRaises(NormalizationError):
                normalize_mt5_import(
                    imported, broker_time_profile_id="roboforex-eet-eest-v0.1"
                )


if __name__ == "__main__":
    unittest.main()
