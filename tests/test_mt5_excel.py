from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from trading_research_lab.mt5_excel import ImportValidationError, import_mt5_excel_report


HEADERS = (
    "Time", "Deal", "Symbol", "Type", "Direction", "Volume", "Price", "Order",
    "Commission", "Swap", "Profit", "Balance", "Comment",
)


def write_report(path: Path, price: float = 158.425, price_format: str = "0.000") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    settings = {
        "Expert:": "EA_DCA_CENT_V1",
        "Symbol:": "USDJPY",
        "Period:": "H4 (2025.01.01 - 2026.01.30)",
        "Currency:": "USD",
        "Initial Deposit:": 15000,
        "Leverage:": "1:1000",
    }
    row = 1
    for key, value in settings.items():
        sheet.cell(row=row, column=1, value=key)
        sheet.cell(row=row, column=4, value=value)
        row += 1
    sheet.cell(row=row, column=1, value="Results")
    row += 1
    sheet.cell(row=row, column=1, value="Deals")
    row += 1
    for column, header in enumerate(HEADERS, 1):
        sheet.cell(row=row, column=column, value=header)
    row += 1
    sheet.append(["2025.01.01 00:00:00", 1, None, "balance", None, None, None, None, 0, 0, 15000, 15000, None])
    sheet.append(["2025.01.10 08:00:00", 2, "USDJPY", "buy", "in", 0.01, price, 2, 0, 0, 0, 15000, "test"])
    sheet.append(["2025.01.10 16:00:00", 3, "USDJPY", "sell", "out", 0.01, 158.552, 3, -0.10, -0.20, 0.80, 15000.50, None])
    for cell in (sheet.cell(row=row + 1, column=7), sheet.cell(row=row + 2, column=7)):
        cell.number_format = price_format
    workbook.save(path)


class Mt5ExcelImportTests(unittest.TestCase):
    def test_preserves_signed_cash_components_and_jpy_scale(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "USDJPY.xlsx"
            write_report(path)

            result = import_mt5_excel_report(path)

        self.assertEqual(result.settings["Symbol"], "USDJPY")
        self.assertEqual(result.opening_balance, 15000)
        self.assertEqual(len(result.deals), 3)
        self.assertEqual(len(result.trading_deals), 2)
        closed_deal = result.deals[-1]
        self.assertEqual(str(closed_deal.price), "158.552")
        self.assertEqual(closed_deal.source_price_scale, 3)
        self.assertEqual(closed_deal.net_cash_change, Decimal("0.50"))
        self.assertEqual(result.net_trading_cash_change, Decimal("0.50"))

    def test_rejects_unsupported_deals_header(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "invalid.xlsx"
            write_report(path)
            workbook = load_workbook(path)
            sheet = workbook.active
            deals_row = next(cell.row for cell in sheet["A"] if cell.value == "Deals")
            sheet.cell(row=deals_row + 1, column=1, value="Timestamp")
            workbook.save(path)

            with self.assertRaises(ImportValidationError):
                import_mt5_excel_report(path)


if __name__ == "__main__":
    unittest.main()
